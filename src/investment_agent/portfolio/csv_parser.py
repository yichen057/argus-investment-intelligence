from __future__ import annotations

import csv
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from typing import Iterable, Sequence
from zipfile import BadZipFile

import xlrd
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from xlrd.biffh import XLRDError


REQUIRED_COLUMNS = {
    "symbol",
    "name",
    "asset_class",
    "quantity",
    "price",
    "market_value",
    "cost_basis",
    "account",
}
SUPPORTED_HOLDINGS_SUFFIXES = {".csv", ".xlsx", ".xls"}
MAX_HOLDINGS_ROWS = 10_000


@dataclass(frozen=True)
class HoldingRow:
    symbol: str
    name: str
    asset_class: str
    quantity: float
    price: float
    market_value: float
    cost_basis: float | None
    account: str | None
    as_of_date: date | None


def parse_holdings_csv(path: Path) -> list[HoldingRow]:
    """Parse a holdings file from a path.

    The historical function name remains for compatibility with evals and scripts,
    but the implementation now accepts CSV, XLSX, and legacy XLS files.
    """

    return parse_holdings_bytes(filename=path.name, raw_content=path.read_bytes())


def parse_holdings_bytes(*, filename: str, raw_content: bytes) -> list[HoldingRow]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_HOLDINGS_SUFFIXES:
        supported = ", ".join(sorted(SUPPORTED_HOLDINGS_SUFFIXES))
        raise ValueError(f"Holdings file must use one of: {supported}.")
    if not raw_content:
        raise ValueError("Holdings file is empty.")

    if suffix == ".csv":
        text = raw_content.decode("utf-8-sig")
        reader = csv.reader(StringIO(text, newline=""))
        return _parse_table(reader, source_label="Holdings CSV")
    if suffix == ".xlsx":
        return _parse_xlsx(raw_content)
    return _parse_xls(raw_content)


def _parse_xlsx(raw_content: bytes) -> list[HoldingRow]:
    try:
        workbook = load_workbook(
            BytesIO(raw_content),
            read_only=True,
            data_only=True,
        )
    except (InvalidFileException, BadZipFile, OSError, ValueError, EOFError) as exc:
        raise ValueError("Holdings XLSX could not be read as an Excel workbook.") from exc

    try:
        worksheet = workbook.active
        return _parse_table(
            worksheet.iter_rows(values_only=True),
            source_label="Holdings XLSX first worksheet",
        )
    finally:
        workbook.close()


def _parse_xls(raw_content: bytes) -> list[HoldingRow]:
    try:
        workbook = xlrd.open_workbook(file_contents=raw_content, on_demand=True)
    except XLRDError as exc:
        raise ValueError("Holdings XLS could not be read as an Excel workbook.") from exc

    try:
        if workbook.nsheets == 0:
            raise ValueError("Holdings XLS is empty.")
        worksheet = workbook.sheet_by_index(0)
        rows = (worksheet.row_values(index) for index in range(worksheet.nrows))
        return _parse_table(rows, source_label="Holdings XLS first worksheet")
    finally:
        workbook.release_resources()


def _parse_table(
    rows: Iterable[Sequence[object]],
    *,
    source_label: str,
) -> list[HoldingRow]:
    iterator = iter(rows)
    try:
        raw_headers = next(iterator)
    except StopIteration as exc:
        raise ValueError(f"{source_label} is empty.") from exc

    headers = [_normalize_key(_cell_text(value)) for value in raw_headers]
    if not any(headers):
        raise ValueError(f"{source_label} is empty.")
    if len({header for header in headers if header}) != len(
        [header for header in headers if header]
    ):
        raise ValueError(f"{source_label} has duplicate columns.")

    normalized_fields = {header for header in headers if header}
    missing = sorted(REQUIRED_COLUMNS - normalized_fields)
    if missing:
        raise ValueError(
            f"{source_label} is missing columns: {', '.join(missing)}"
        )

    parsed_rows: list[HoldingRow] = []
    for row_number, raw_row in enumerate(iterator, start=2):
        if row_number > MAX_HOLDINGS_ROWS + 1:
            raise ValueError(
                f"{source_label} exceeds the {MAX_HOLDINGS_ROWS:,} position limit."
            )
        values = [_cell_text(value) for value in raw_row]
        if not any(values):
            continue
        row = {
            header: values[index] if index < len(values) else ""
            for index, header in enumerate(headers)
            if header
        }
        parsed_rows.append(_parse_row(row, row_number=row_number))

    if not parsed_rows:
        raise ValueError(f"{source_label} has no position rows.")
    as_of_dates = {row.as_of_date for row in parsed_rows if row.as_of_date is not None}
    if len(as_of_dates) > 1:
        raise ValueError(
            f"{source_label} contains multiple as_of_date values; "
            "one import must represent one portfolio snapshot."
        )
    return parsed_rows


def _parse_row(row: dict[str, str], *, row_number: int) -> HoldingRow:
    normalized = {_normalize_key(key): (value or "").strip() for key, value in row.items()}
    symbol = normalized["symbol"].upper()
    name = normalized["name"]
    asset_class = normalized["asset_class"]
    if not symbol:
        raise ValueError(f"Row {row_number}: symbol is required.")
    if not name:
        raise ValueError(f"Row {row_number}: name is required.")
    if not asset_class:
        raise ValueError(f"Row {row_number}: asset_class is required.")

    quantity = _parse_float(normalized["quantity"], row_number=row_number, field="quantity")
    price = _parse_float(normalized["price"], row_number=row_number, field="price")
    market_value = _parse_float(
        normalized["market_value"], row_number=row_number, field="market_value"
    )
    cost_basis = _parse_optional_float(
        normalized["cost_basis"], row_number=row_number, field="cost_basis"
    )

    return HoldingRow(
        symbol=symbol,
        name=name,
        asset_class=asset_class,
        quantity=quantity,
        price=price,
        market_value=market_value,
        cost_basis=cost_basis,
        account=normalized["account"] or None,
        as_of_date=_parse_optional_date(
            normalized.get("as_of_date", ""),
            row_number=row_number,
        ),
    )


def _normalize_key(value: str) -> str:
    return value.strip().lower().replace(" ", "_")


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _parse_optional_float(value: str, *, row_number: int, field: str) -> float | None:
    if not value:
        return None
    return _parse_float(value, row_number=row_number, field=field)


def _parse_optional_date(value: str, *, row_number: int) -> date | None:
    if not value:
        return None
    normalized = value.strip()
    for candidate in (normalized, normalized[:10]):
        try:
            return date.fromisoformat(candidate)
        except ValueError:
            continue
    for date_format in ("%m/%d/%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(normalized, date_format).date()
        except ValueError:
            continue
    raise ValueError(
        f"Row {row_number}: as_of_date must use YYYY-MM-DD or MM/DD/YYYY."
    )


def _parse_float(value: str, *, row_number: int, field: str) -> float:
    try:
        parsed = float(value.replace("$", "").replace(",", "").strip())
    except ValueError as exc:
        raise ValueError(f"Row {row_number}: {field} must be numeric.") from exc
    if parsed < 0:
        raise ValueError(f"Row {row_number}: {field} cannot be negative.")
    return parsed
